"""
diagnose_pptx.py — 判断 PPT 里的图是"原生矢量(形状+连接线)"还是"贴进去的位图"
在公司电脑本地跑即可,不需要把文件传出去。
用法: python diagnose_pptx.py 你的文件.pptx
依赖: pip install python-pptx lxml
"""
import sys
from pptx import Presentation
from lxml import etree


def ln(el):                       # 取 XML 标签的本名,如 sp / pic / cxnSp
    return etree.QName(el).localname


def walk(container, id2label, edges, counts):
    """递归遍历(含组合 grpSp),统计形状类型并抽取连接线的两端"""
    for el in container:
        tag = ln(el)
        if tag == 'grpSp':                         # 组合,递归进去
            walk(el, id2label, edges, counts)
            continue
        if tag not in ('sp', 'pic', 'cxnSp', 'graphicFrame'):
            continue
        counts[tag] = counts.get(tag, 0) + 1
        cNvPr = next((c for c in el.iter() if ln(c) == 'cNvPr'), None)
        sid = cNvPr.get('id') if cNvPr is not None else None
        name = cNvPr.get('name', '') if cNvPr is not None else ''
        text = ''.join((t.text or '') for t in el.iter() if ln(t) == 't').strip()
        if sid is not None:
            id2label[sid] = (text or name).replace('\n', ' ')[:40] or f'#{sid}'
        if tag == 'cxnSp':                          # 连接线:读起点/终点连到哪个形状
            st = next((c for c in el.iter() if ln(c) == 'stCxn'), None)
            en = next((c for c in el.iter() if ln(c) == 'endCxn'), None)
            edges.append((st.get('id') if st is not None else None,
                          en.get('id') if en is not None else None))


def diagnose(path):
    prs = Presentation(path)
    for i, slide in enumerate(prs.slides, 1):
        id2label, edges, counts = {}, [], {}
        walk(slide.shapes._spTree, id2label, edges, counts)
        pics = counts.get('pic', 0)
        conns = counts.get('cxnSp', 0)
        shapes = counts.get('sp', 0)
        frames = counts.get('graphicFrame', 0)   # 表格 / SmartArt

        if conns > 0:
            verdict = '矢量 · 有连接线(骨架可读)✅'
        elif shapes > pics and shapes >= 3:
            verdict = '矢量 · 形状拼的(方向可能靠位置,无显式连接线)'
        elif pics >= 1 and shapes <= 2:
            verdict = '位图为主 · 很可能是贴图 ❌(连接线路走不通)'
        else:
            verdict = '不确定 / 可能是 SmartArt 或表格'

        print(f'\n=== 第 {i} 页 ===')
        print(f'图片:{pics}  连接线:{conns}  形状:{shapes}  表格/SmartArt:{frames}')
        print(f'判定:{verdict}')
        if edges:
            print('连接线 起点 → 终点(能显示名字=方向确实可读):')
            for s, e in edges:
                print(f'   {id2label.get(s, s)}  →  {id2label.get(e, e)}')


if __name__ == '__main__':
    diagnose(sys.argv[1])


