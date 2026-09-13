"""
flowchart_graph.py — 从 PPT 矢量流程图重建"有向图"(节点+带方向的边)
只处理原生矢量(形状+连接线);贴图走你现有 OCR,SmartArt 另处理。
输出可存成 artifact JSON,路径挂到 chunk metadata。
"""
from pptx import Presentation
from lxml import etree

A = 'http://schemas.openxmlformats.org/drawingml/2006/main'

def ln(el):
    return etree.QName(el).localname

def _text(el):
    return ''.join((t.text or '') for t in el.iter() if ln(t) == 't').strip().replace('\n', ' ')

def _bbox(el):
    """取形状的 off(x,y) 和 ext(cx,cy);拿不到返回 None"""
    xfrm = next((c for c in el.iter() if ln(c) == 'xfrm'), None)
    if xfrm is None:
        return None
    off = next((c for c in xfrm if ln(c) == 'off'), None)
    ext = next((c for c in xfrm if ln(c) == 'ext'), None)
    if off is None or ext is None:
        return None
    x, y = int(off.get('x', 0)), int(off.get('y', 0))
    cx, cy = int(ext.get('cx', 0)), int(ext.get('cy', 0))
    return (x, y, cx, cy)

def _center(b):
    return (b[0] + b[2] / 2, b[1] + b[3] / 2)

def _collect(container, nodes, connectors):
    """递归收集:节点(sp)、连接线(cxnSp);进 group 递归"""
    for el in container:
        tag = ln(el)
        if tag == 'grpSp':
            _collect(el, nodes, connectors)
            continue
        if tag not in ('sp', 'cxnSp'):
            continue
        cNvPr = next((c for c in el.iter() if ln(c) == 'cNvPr'), None)
        sid = cNvPr.get('id') if cNvPr is not None else None
        name = cNvPr.get('name', '') if cNvPr is not None else ''
        if tag == 'sp':
            nodes[sid] = {'id': sid, 'label': _text(el) or '', 'name': name, 'bbox': _bbox(el)}
        else:  # cxnSp 连接线
            st = next((c for c in el.iter() if ln(c) == 'stCxn'), None)
            en = next((c for c in el.iter() if ln(c) == 'endCxn'), None)
            ln_el = next((c for c in el.iter() if ln(c) == 'ln'), None)
            head = tail = None
            if ln_el is not None:
                h = next((c for c in ln_el if ln(c) == 'headEnd'), None)
                t = next((c for c in ln_el if ln(c) == 'tailEnd'), None)
                head = h.get('type') if h is not None else None
                tail = t.get('type') if t is not None else None
            connectors.append({
                'start_id': st.get('id') if st is not None else None,
                'end_id': en.get('id') if en is not None else None,
                'bbox': _bbox(el), 'head': head, 'tail': tail,
            })

def _nearest_node(point, nodes):
    """几何兜底:端点坐标 -> 最近的有 bbox 的节点 id"""
    best, best_d = None, None
    for n in nodes.values():
        if not n['bbox']:
            continue
        cx, cy = _center(n['bbox'])
        d = (cx - point[0]) ** 2 + (cy - point[1]) ** 2
        if best_d is None or d < best_d:
            best, best_d = n['id'], d
    return best

def _borrow_labels(nodes):
    """给没文字的形状借标签:找 bbox 落在它内部、且有文字的其它节点"""
    boxed = [n for n in nodes.values() if n['bbox']]
    for n in nodes.values():
        if n['label'] or not n['bbox']:
            continue
        x, y, cx, cy = n['bbox']
        for m in boxed:
            if m is n or not m['label']:
                continue
            mcx, mcy = _center(m['bbox'])
            if x <= mcx <= x + cx and y <= mcy <= y + cy:
                n['label'] = m['label']
                break

def build_graph(slide):
    nodes, connectors = {}, []
    _collect(slide.shapes._spTree, nodes, connectors)
    _borrow_labels(nodes)

    edges = []
    for c in connectors:
        s, e = c['start_id'], c['end_id']
        # 几何兜底:某端没连上形状,就按连接线两端坐标找最近节点
        if (s is None or e is None) and c['bbox']:
            x, y, cx, cy = c['bbox']
            p1, p2 = (x, y), (x + cx, y + cy)
            s = s or _nearest_node(p1, nodes)
            e = e or _nearest_node(p2, nodes)
        if s is None and e is None:
            continue  # 纯噪声,丢
        # 方向:箭头在 tail 端 => start->end;在 head 端 => 反过来
        directed = True
        arrow = lambda v: v and v.lower() != 'none'
        if arrow(c['tail']) and not arrow(c['head']):
            src, dst = s, e
        elif arrow(c['head']) and not arrow(c['tail']):
            src, dst = e, s
        else:
            src, dst, directed = s, e, False  # 双箭头/无箭头:方向存疑
        label = lambda i: (nodes.get(i, {}).get('label') or nodes.get(i, {}).get('name') or f'#{i}')
        edges.append({'src': label(src), 'dst': label(dst), 'directed': directed})

    return {
        'nodes': [{'id': n['id'], 'label': n['label'] or n['name']} for n in nodes.values()],
        'edges': edges,
        'classification': 'vector' if connectors else 'sparse',
    }

def to_mermaid(graph):
    lines = ['flowchart TD']
    for e in graph['edges']:
        arrow = '-->' if e['directed'] else '---'
        lines.append(f'    ["{e["src"]}"] {arrow} ["{e["dst"]}"]')
    return '\n'.join(lines)


if __name__ == '__main__':
    import sys, json
    prs = Presentation(sys.argv[1])
    for i, slide in enumerate(prs.slides, 1):
        g = build_graph(slide)
        if g['edges']:
            print(f'\n=== 第 {i} 页 · {len(g["edges"])} 条边 ===')
            print(json.dumps(g, ensure_ascii=False, indent=2))
            print(to_mermaid(g))
</parameter>



